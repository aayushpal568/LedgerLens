import csv
import io
import os
import socket
import subprocess
import time
from pathlib import Path

import requests
from openpyxl import load_workbook
from pypdf import PdfReader

APP = Path(os.environ["LOCALAPPDATA"]) / "LedgerLens" / "ledgerlens.exe"
SCAN_ID = "22d6e556-9d6e-4cc9-9cb1-72e15d48c95b"
OUT = Path(__file__).resolve().parent / "final_ui_downloads"
OUT.mkdir(exist_ok=True)


def kill_all():
    subprocess.run([
        "powershell.exe", "-NoProfile", "-Command",
        "Get-Process ledgerlens,ledgerlens-backend,tauri-driver,msedgedriver -ErrorAction SilentlyContinue | Stop-Process -Force",
    ], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def wait_backend(seconds=60):
    start = time.time()
    while time.time() - start < seconds:
        try:
            if requests.get("http://127.0.0.1:8001/api/firm", timeout=2).status_code == 200:
                return True, round(time.time() - start, 1)
        except requests.RequestException:
            pass
        time.sleep(1)
    return False, round(time.time() - start, 1)


def process_count():
    out = subprocess.run([
        "powershell.exe", "-NoProfile", "-Command",
        "@(Get-Process ledgerlens,ledgerlens-backend,tauri-driver,msedgedriver -ErrorAction SilentlyContinue).Count",
    ], capture_output=True, text=True).stdout.strip()
    return int(out or 0)

results = []
def record(name, passed, detail=""):
    results.append((name, passed, detail))
    print(f"[{'PASS' if passed else 'FAIL'}] {name} - {detail}")

kill_all()
proc = subprocess.Popen([str(APP)])
ready, elapsed = wait_backend()
record("Backend available for report validation", ready, f"{elapsed}s")

if ready:
    for fmt in ("csv", "xlsx", "pdf"):
        response = requests.get(f"http://127.0.0.1:8001/api/scans/{SCAN_ID}/report?format={fmt}", timeout=60)
        target = OUT / f"final_report.{fmt}"
        target.write_bytes(response.content)
        record(f"{fmt.upper()} report physically readable", response.ok and target.stat().st_size > 0, f"{target.stat().st_size} bytes")

    csv_path = OUT / "final_report.csv"
    rows = list(csv.reader(io.StringIO(csv_path.read_text(encoding="utf-8-sig"))))
    dangerous = [cell for row in rows[1:] for cell in row if cell.startswith(("=", "+", "-", "@"))]
    record("CSV formula-like cells neutralized", not dangerous, repr(dangerous[:3]))

    wb = load_workbook(OUT / "final_report.xlsx", data_only=False)
    dangerous_xlsx = []
    for row in wb.active.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                dangerous_xlsx.append(cell.value)
    record("XLSX formula-like cells neutralized", not dangerous_xlsx, repr(dangerous_xlsx[:3]))

    pdf = PdfReader(str(OUT / "final_report.pdf"))
    text = "\n".join((p.extract_text() or "") for p in pdf.pages)
    record("PDF parses after special-character notes", len(pdf.pages) > 0 and "Document Review Report" in text, f"{len(pdf.pages)} pages")

kill_all(); time.sleep(2)
record("No orphan processes after report session", process_count() == 0, str(process_count()))

cold_ok = 0
for cycle in range(1, 4):
    kill_all(); time.sleep(1)
    subprocess.Popen([str(APP)])
    ready, elapsed = wait_backend(60)
    kill_all(); time.sleep(2)
    orphans = process_count()
    passed = ready and orphans == 0
    cold_ok += int(passed)
    record(f"Cold-start cycle {cycle}", passed, f"backend={ready} in {elapsed}s, orphans={orphans}")
record("Three cold-start cycles", cold_ok == 3, f"{cold_ok}/3")

raise SystemExit(0 if all(p for _, p, _ in results) else 1)
