import csv
import sys
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

paths = [Path(p) for p in sys.argv[1:]]
assert len(paths) == 3
csv_path = next(p for p in paths if p.suffix.lower() == ".csv")
xlsx_path = next(p for p in paths if p.suffix.lower() == ".xlsx")
pdf_path = next(p for p in paths if p.suffix.lower() == ".pdf")

rows = list(csv.reader(csv_path.read_text(encoding="utf-8-sig").splitlines()))
assert len(rows) >= 2 and rows[0][0] == "Category"
assert not [v for row in rows[1:] for v in row if v.startswith(("=", "+", "-", "@"))]

wb = load_workbook(xlsx_path, data_only=False)
assert wb.active.title == "Exceptions" and wb.active.max_row >= 2
assert not [c.value for row in wb.active.iter_rows(min_row=2) for c in row
            if isinstance(c.value, str) and c.value.startswith(("=", "+", "-", "@"))]

pdf = PdfReader(str(pdf_path))
text = "\n".join(page.extract_text() or "" for page in pdf.pages)
assert pdf.pages and "Document Review Report" in text
print(f"VALID CSV={csv_path.stat().st_size} XLSX={xlsx_path.stat().st_size} PDF={pdf_path.stat().st_size}")
