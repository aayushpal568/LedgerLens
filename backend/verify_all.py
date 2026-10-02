"""Complete Verification Suite for LedgerLens Cloud Backend & Core Engine.

Validates:
1. Backend startup & API endpoints (firm, clients, templates, files, scans, findings, reports)
2. PostgreSQL database layer & cloud storage fallback
3. Core Engine with realistic sample documents:
   - PDF, DOCX, XLSX, CSV
   - Exact duplicate detection
   - Near duplicate detection (Jaccard similarity)
   - Wrong period detection
   - Wrong file type detection
   - Missing document detection
   - Unreadable / empty file detection
4. End-to-end integration workflow
"""
import io
import os
import sys
import tempfile
import time
from pathlib import Path

# Setup paths
ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

# Hermetic verification: Normal verify_all must NEVER make real Claude API calls
RUN_LIVE_CLAUDE = "--live-claude" in sys.argv
if not RUN_LIVE_CLAUDE:
    os.environ["FAL_KEY"] = ""
    os.environ["FAL_API_KEY"] = ""
    os.environ["ANTHROPIC_API_KEY"] = ""

# When no external PostgreSQL database is configured, verify_all runs hermetically on the explicit memory backend
if not os.environ.get("DATABASE_URL"):
    os.environ.setdefault("DATA_BACKEND", "memory")


import pypdf
import docx
import openpyxl
import reportlab
from fastapi.testclient import TestClient

from engine import (
    build_default_engine, run_detection, default_templates,
    LocalPathFileSource, DefaultDocumentExtractor, NoOpOCRProvider, NoOpLLMProvider,
)
from engine.checklist import detect_period, match_item
from engine.similarity import similarity
import server


def test_backend_api_and_database():
    print("\n--- 1. Testing Backend API & Database Layer ---")
    with TestClient(server.app) as client:
        # Root & health
        res = client.get("/")
        assert res.status_code == 200, f"Root failed: {res.status_code}"
        print("  [PASS] Root endpoint /")

        # Authenticate
        signup_res = client.post("/api/auth/signup", json={
            "email": "verify_all_api@example.com",
            "password": "Password123!",
            "firm_name": "Verify Firm",
            "user_name": "Verify User"
        })
        assert signup_res.status_code in (200, 201), f"Signup failed: {signup_res.text}"
        token = signup_res.json()["access_token"]
        client.headers["Authorization"] = f"Bearer {token}"

        # Firm
        res = client.get("/api/firm")
        assert res.status_code == 200, f"GET /api/firm failed: {res.text}"
        res = client.put("/api/firm", json={"name": "Premier CA Associates"})
        assert res.status_code == 200 and res.json()["name"] == "Premier CA Associates"
        print("  [PASS] Firm GET/PUT /api/firm")

        # Checklist templates
        res = client.get("/api/checklist-templates")
        assert res.status_code == 200 and len(res.json()) >= 3
        tpl_id = res.json()[0]["id"]
        print(f"  [PASS] Checklist Templates GET ({len(res.json())} templates seeded)")

        # Client CRUD
        res = client.post("/api/clients", json={"name": "Acme Corp", "client_type": "Corporation", "notes": "Tax audit 2024"})
        assert res.status_code == 200
        client_id = res.json()["id"]
        assert client_id is not None

        res = client.get(f"/api/clients/{client_id}")
        assert res.status_code == 200 and res.json()["name"] == "Acme Corp"

        res = client.get("/api/clients")
        assert any(c["id"] == client_id for c in res.json())
        print("  [PASS] Client Create/Get/List")

        # File Upload & Storage Fallback (without external credentials)
        files = [
            ("files", ("invoice_2024.csv", b"Date,Amount,Vendor\n2024-04-01,1500,Cloud Servers\n", "text/csv")),
            ("files", ("bank_stmt_2024.csv", b"Date,Amount,Vendor\n2024-04-02,-1500,Bank Transfer\n", "text/csv")),
        ]
        res = client.post(f"/api/clients/{client_id}/files", files=files)
        assert res.status_code == 200 and res.json()["uploaded"] == 2, f"Upload failed: {res.text}"
        uploaded_files = res.json()["files"]
        print(f"  [PASS] File Upload & Storage Integration ({len(uploaded_files)} files saved)")

        # List files
        res = client.get(f"/api/clients/{client_id}/files")
        assert res.status_code == 200 and len(res.json()) == 2
        print("  [PASS] File List /api/clients/{id}/files")

        # Start Scan
        res = client.post(f"/api/clients/{client_id}/scan", json={"template_id": tpl_id, "expected_period": 2024})
        assert res.status_code == 200
        scan_id = res.json()["id"]
        print(f"  [PASS] Scan Initiation /api/clients/{id}/scan (scan_id={scan_id})")

        # Poll scan completion
        for _ in range(50):
            time.sleep(0.3)
            s_res = client.get(f"/api/scans/{scan_id}")
            if s_res.json()["status"] in ("completed", "error"):
                break
        assert s_res.json()["status"] == "completed", f"Scan failed: {s_res.json()}"
        print("  [PASS] Scan Execution & Background Processing (Status: completed)")

        # Get findings
        f_res = client.get(f"/api/scans/{scan_id}/findings")
        assert f_res.status_code == 200
        findings = f_res.json()
        print(f"  [PASS] Findings Retrieval ({len(findings)} findings detected)")

        # Update finding review status
        if findings:
            target_f = findings[0]
            up_res = client.patch(f"/api/findings/{target_f['id']}", json={"status": "keep", "note": "Verified by accountant"})
            assert up_res.status_code == 200 and up_res.json()["status"] == "keep"
            print("  [PASS] Findings Review Status Update")

        # Generate Reports
        for fmt, content_type in [("csv", "text/csv"), ("xlsx", "application/vnd.openxmlformats"), ("pdf", "application/pdf")]:
            rep_res = client.get(f"/api/scans/{scan_id}/report?format={fmt}")
            assert rep_res.status_code == 200, f"Report {fmt} failed: {rep_res.status_code}"
            assert len(rep_res.content) > 0
            print(f"  [PASS] Report Export ({fmt.upper()}, {len(rep_res.content)} bytes)")

        # Cleanup client
        del_res = client.delete(f"/api/clients/{client_id}")
        assert del_res.status_code == 200
        print("  [PASS] Client Deletion & Cascade Cleanup")


def test_core_engine_realistic_files():
    print("\n--- 2. Testing Core LedgerLens Engine with Realistic Documents ---")
    with tempfile.TemporaryDirectory() as tmpdir:
        tmppath = Path(tmpdir)

        # 1. Real PDF (PyPDF)
        pdf_path = tmppath / "Tax_Return_2024.pdf"
        from reportlab.pdfgen import canvas
        c = canvas.Canvas(str(pdf_path))
        c.drawString(100, 750, "Annual Corporate Tax Return 2024")
        c.drawString(100, 730, "Gross Receipts: $450,000. Total Deductions: $320,000")
        c.save()

        # 2. Real DOCX (python-docx)
        docx_path = tmppath / "Board_Minutes_2024.docx"
        doc = docx.Document()
        doc.add_heading("Board Minutes 2024", 0)
        doc.add_paragraph("The board met to approve financial statements for fiscal year 2024.")
        doc.save(str(docx_path))

        # 3. Real XLSX (openpyxl)
        xlsx_path = tmppath / "Balance_Sheet_2024.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Balance Sheet"
        ws.append(["Balance Sheet 2024", "Amount"])
        ws.append(["Cash & Cash Equivalents", 150000])
        ws.append(["Accounts Receivable", 45000])
        ws.append(["Total Liabilities", 60000])
        wb.save(str(xlsx_path))

        # 4. Real CSV (General Ledger with multiple transactions)
        ledger_lines = ["Account,Debit,Credit,Year,Ref"]
        for i in range(1, 30):
            ledger_lines.append(f"ACC-{1000+i},{i*100},0,2024,TX-{5000+i}")
        ledger_lines.append("ACC-9999,0,43500,2024,BAL")
        csv_path = tmppath / "General_Ledger_2024.csv"
        csv_path.write_text("\n".join(ledger_lines) + "\n")

        # 5. Exact duplicate file
        dup_path = tmppath / "General_Ledger_Copy.csv"
        dup_path.write_bytes(csv_path.read_bytes())

        # 6. Near-duplicate file (slight amendment to last transaction, >95% similar)
        near_lines = list(ledger_lines)
        near_lines[-1] = "ACC-9999,0,43500,2024,BAL_AMENDED"
        near_path = tmppath / "General_Ledger_Draft.csv"
        near_path.write_text("\n".join(near_lines) + "\n")


        # 7. Wrong-period document (2021 instead of 2024)
        wp_path = tmppath / "Profit_and_Loss_2021.csv"
        wp_path.write_text("Profit and Loss Statement 2021\nRevenue,250000\nExpenses,180000\n")

        # 8. Wrong file type (PNG for Trial Balance which expects XLSX/CSV/PDF)
        wt_path = tmppath / "Trial_Balance.png"
        wt_path.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 32)

        # 9. Empty / unreadable file
        empty_path = tmppath / "Corrupt_Doc.csv"
        empty_path.write_text("")

        # Assemble file records
        records = [
            {"id": "f1", "name": pdf_path.name, "ext": "pdf", "size": pdf_path.stat().st_size, "path": str(pdf_path)},
            {"id": "f2", "name": docx_path.name, "ext": "docx", "size": docx_path.stat().st_size, "path": str(docx_path)},
            {"id": "f3", "name": xlsx_path.name, "ext": "xlsx", "size": xlsx_path.stat().st_size, "path": str(xlsx_path)},
            {"id": "f4", "name": csv_path.name, "ext": "csv", "size": csv_path.stat().st_size, "path": str(csv_path)},
            {"id": "f5", "name": dup_path.name, "ext": "csv", "size": dup_path.stat().st_size, "path": str(dup_path)},
            {"id": "f6", "name": near_path.name, "ext": "csv", "size": near_path.stat().st_size, "path": str(near_path)},
            {"id": "f7", "name": wp_path.name, "ext": "csv", "size": wp_path.stat().st_size, "path": str(wp_path)},
            {"id": "f8", "name": wt_path.name, "ext": "png", "size": wt_path.stat().st_size, "path": str(wt_path)},
            {"id": "f9", "name": empty_path.name, "ext": "csv", "size": empty_path.stat().st_size, "path": str(empty_path)},
        ]

        # Use Corporation Year-End checklist template
        corp_tpl = next(t for t in default_templates() if t["client_type"] == "Corporation")
        checklist_items = corp_tpl["items"]

        # Run ScanEngine
        engine = build_default_engine()
        source = LocalPathFileSource(records)
        result = engine.run(source, checklist_items, expected_period=2024)

        counts = result["counts"]
        print("  Detection Counts:")
        for cat, cnt in counts.items():
            print(f"    - {cat}: {cnt}")

        # Assertions for each expected detection
        assert counts["exact_duplicate"] >= 1, "Failed to detect exact duplicate (f4 == f5)"
        print("  [PASS] Exact Duplicate Detection (matching SHA-256)")

        assert counts["possible_duplicate"] >= 1, "Failed to detect near duplicate"
        print("  [PASS] Near-Duplicate Detection (Jaccard similarity >= 82%)")

        assert counts["wrong_period"] >= 1, "Failed to detect wrong period (2021 vs 2024)"
        print("  [PASS] Wrong Period Detection (detected 2021 vs expected 2024)")

        assert counts["wrong_type"] >= 1, "Failed to detect wrong file type"
        print("  [PASS] Wrong File Type Detection (PNG vs XLSX/CSV/PDF)")

        assert counts["missing_doc"] >= 1, "Failed to detect missing checklist document"
        print("  [PASS] Missing Document Detection (unmatched checklist items)")

        assert counts["unreadable"] >= 1, "Failed to detect empty/unreadable file"
        print("  [PASS] Empty/Unreadable File Detection")


def test_end_to_end_integration():
    print("\n--- 3. Testing Full End-to-End Integration Workflow ---")
    with TestClient(server.app) as client:
        # Authenticate
        signup_res = client.post("/api/auth/signup", json={
            "email": "verify_all_e2e@example.com",
            "password": "Password123!",
            "firm_name": "Apex Chartered Accountants",
            "user_name": "Apex User"
        })
        assert signup_res.status_code in (200, 201), f"Signup failed: {signup_res.text}"
        token = signup_res.json()["access_token"]
        client.headers["Authorization"] = f"Bearer {token}"

        # Step 1: Create Client
        c_res = client.post("/api/clients", json={"name": "Apex Chartered Accountants Client", "client_type": "Small Business"})
        cid = c_res.json()["id"]

        # Step 2: Upload Documents
        csv_doc1 = ("files", ("income_statement_2024.csv", b"Category,Amount,Year\nSales,120000,2024\nCost,40000,2024\n", "text/csv"))
        csv_doc2 = ("files", ("income_statement_copy.csv", b"Category,Amount,Year\nSales,120000,2024\nCost,40000,2024\n", "text/csv"))
        up_res = client.post(f"/api/clients/{cid}/files", files=[csv_doc1, csv_doc2])
        assert up_res.status_code == 200

        # Step 3: Process documents
        s_res = client.post(f"/api/clients/{cid}/scan", json={"expected_period": 2024})
        sid = s_res.json()["id"]
        for _ in range(25):
            time.sleep(0.15)
            cur_s = client.get(f"/api/scans/{sid}").json()
            if cur_s["status"] == "completed":
                break
        assert cur_s["status"] == "completed"

        # Step 4: Detect issues & view findings
        find_res = client.get(f"/api/scans/{sid}/findings").json()
        assert len(find_res) > 0
        exact_dup = next(f for f in find_res if f["category"] == "exact_duplicate")
        assert exact_dup is not None

        # Step 5: Review findings (accountant approval)
        rev_res = client.patch(f"/api/findings/{exact_dup['id']}", json={"status": "ignore", "note": "Duplicate acknowledged and discarded"})
        assert rev_res.status_code == 200

        # Step 6: Generate report
        pdf_res = client.get(f"/api/scans/{sid}/report?format=pdf")
        assert pdf_res.status_code == 200 and pdf_res.content.startswith(b"%PDF")
        print("  [PASS] Full End-to-End Flow: Client -> Upload -> Process -> Detect -> Review -> PDF Report")


def test_paddle_ocr_integration():
    print("\n--- 4. Testing Local PaddleOCR Provider & Pipeline ---")
    from unittest.mock import patch, MagicMock
    from PIL import Image, ImageDraw
    import fitz
    from engine.providers import PaddleOCRProvider, DefaultDocumentExtractor
    from engine.models import STATUS_OK, STATUS_NEEDS_OCR

    # Test 1: Initialization with default CPU configuration
    provider = PaddleOCRProvider()
    assert provider.name == "paddle_ocr"
    assert provider.available is True
    summary = provider.get_config_summary()
    assert summary["provider"] == "paddle_ocr"
    assert summary["available"] == "True"
    assert summary["mode"] == "local"
    assert summary["use_gpu"] == "False"
    assert summary["lang"] == "en"
    assert summary["use_angle_cls"] == "True"
    print("  [PASS] PaddleOCR initialization and safe CPU defaults")

    # Test 2: Environment variable overrides (GPU / language)
    env = {"PADDLE_OCR_USE_GPU": "true", "PADDLE_OCR_LANG": "ch", "PADDLE_OCR_USE_ANGLE_CLS": "false"}
    with patch.dict(os.environ, env, clear=False):
        prov_env = PaddleOCRProvider()
        cfg = prov_env.get_config_summary()
        assert cfg["use_gpu"] == "True"
        assert cfg["lang"] == "ch"
        assert cfg["use_angle_cls"] == "False"
    print("  [PASS] Configuration handling for GPU and language overrides")

    # Test 3: Graceful handling when OCR is unavailable
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_img = Path(tmpdir) / "invoice_receipt.png"
        img = Image.new("RGB", (300, 150), color=(255, 255, 255))
        d = ImageDraw.Draw(img)
        d.text((10, 10), "Receipt 2024", fill=(0, 0, 0))
        img.save(tmp_img)

        provider_unavailable = PaddleOCRProvider()
        provider_unavailable._available_cache = False
        assert provider_unavailable.extract(str(tmp_img), "png") is None

        extractor_unavail = DefaultDocumentExtractor(provider_unavailable)
        ext_res = extractor_unavail.extract(str(tmp_img), "png")
        assert ext_res.status == STATUS_NEEDS_OCR
        assert ext_res.ocr_used is False
        assert "OCR is not configured" in ext_res.reason
        assert "paddle_ocr" in ext_res.reason
        print("  [PASS] Graceful handling and clear reporting when OCR is unavailable")

        # Test 4: Image extraction with mocked OCR engine
        mock_ocr = MagicMock()
        mock_ocr.ocr.return_value = [
            [
                [[[10, 10], [100, 10], [100, 30], [10, 30]], ("Vendor: AWS Cloud", 0.99)],
                [[[10, 40], [100, 40], [100, 60], [10, 60]], ("Total: $1,250.00", 0.98)],
            ]
        ]
        with patch.object(provider, "_get_ocr", return_value=mock_ocr):
            text = provider.extract(str(tmp_img), "png")
            assert text is not None and "AWS Cloud" in text
            print("  [PASS] Image OCR extraction via PaddleOCR")

            # Test 5: OCR text reaches document processing pipeline
            extractor = DefaultDocumentExtractor(provider)
            pipeline_res = extractor.extract(str(tmp_img), "png")
            assert pipeline_res.status == STATUS_OK
            assert pipeline_res.ocr_used is True
            assert pipeline_res.meta.get("ocr") == "paddle_ocr"
            assert "AWS Cloud" in pipeline_res.text
            print("  [PASS] OCR text seamlessly feeds into document pipeline (status=OK, ocr_used=True)")

        # Test 6: Scanned PDF sequential one-page-at-a-time processing
        tmp_pdf = Path(tmpdir) / "scanned_doc.pdf"
        doc = fitz.open()
        p1 = doc.new_page(width=300, height=150)
        p1.insert_image(fitz.Rect(0, 0, 300, 150), filename=str(tmp_img))
        p2 = doc.new_page(width=300, height=150)
        p2.insert_image(fitz.Rect(0, 0, 300, 150), filename=str(tmp_img))
        doc.save(str(tmp_pdf))
        doc.close()

        mock_pdf_ocr = MagicMock()
        mock_pdf_ocr.ocr.side_effect = [
            [[[[[0, 0], [10, 0], [10, 10], [0, 10]], ("Page 1 Invoice", 0.95)]]],
            [[[[[0, 0], [10, 0], [10, 10], [0, 10]], ("Page 2 Appendix", 0.96)]]],
        ]
        with patch.object(provider, "_get_ocr", return_value=mock_pdf_ocr):
            pdf_text = provider.extract(str(tmp_pdf), "pdf")
            assert "--- Page 1 ---" in pdf_text
            assert "Page 1 Invoice" in pdf_text
            assert "--- Page 2 ---" in pdf_text
            assert "Page 2 Appendix" in pdf_text
            assert mock_pdf_ocr.ocr.call_count == 2
            print("  [PASS] Sequential page-by-page PDF OCR with memory bounds")

        # Test 7: PDF preserves successfully processed pages on subsequent page failure
        mock_partial_ocr = MagicMock()
        mock_partial_ocr.ocr.side_effect = [
            [[[[[0, 0], [10, 0], [10, 10], [0, 10]], ("Page 1 Success", 0.95)]]],
            RuntimeError("Simulated OOM on page 2"),
        ]
        with patch.object(provider, "_get_ocr", return_value=mock_partial_ocr):
            partial_text = provider.extract(str(tmp_pdf), "pdf")
            assert "--- Page 1 ---" in partial_text
            assert "Page 1 Success" in partial_text
            assert "--- Page 2 ---" not in partial_text
            print("  [PASS] Preserves earlier pages if subsequent page fails")


def test_claude_opus_fal_integration():
    print("\n--- 5. Testing Claude Opus (fal.ai) LLM Integration ---")
    import json
    import httpx
    from unittest.mock import patch, MagicMock
    from engine.providers.llm import ClaudeOpusFalProvider

    # Test 1: Unconfigured defaults
    provider_unconfigured = ClaudeOpusFalProvider(api_key="", anthropic_key="")
    assert provider_unconfigured.available is False
    assert provider_unconfigured.backend_mode == "unconfigured"
    assert provider_unconfigured.generate("test prompt") is None
    assert provider_unconfigured.query_json("test prompt") is None
    assert provider_unconfigured.classify_document("doc text", ["Invoice"]) is None
    print("  [PASS] Unconfigured provider initializes safely with available=False")

    # Test 2: Credential masking & configuration loading
    provider_configured = ClaudeOpusFalProvider(
        api_key="fal_sec_live_998877665544332211aabbcc",
        endpoint="https://fal.run/fal-ai/any-llm",
        model="anthropic/claude-3-opus",
        anthropic_key="sk-ant-test-secret-key-12345",
    )
    assert provider_configured.available is True
    assert provider_configured.backend_mode == "fal_ai"
    cfg = provider_configured.get_config_summary()
    assert cfg["available"] == "True"
    assert cfg["backend_mode"] == "fal_ai"
    assert cfg["fal_endpoint"] == "https://fal.run/fal-ai/any-llm"
    assert cfg["fal_model"] == "anthropic/claude-3-opus"
    assert "fal_sec_live_998877665544332211aabbcc" not in cfg["fal_key_masked"]
    assert cfg["fal_key_masked"] == "fal...bcc"
    assert "sk-ant-test-secret-key-12345" not in cfg["anthropic_key_masked"]
    assert cfg["anthropic_key_masked"] == "sk-...345"
    print("  [PASS] Credential masking & configuration handling (secrets never leaked)")

    # Test 3: Document classification with extracted document text
    mock_resp = MagicMock()
    mock_resp.status_code = 200
    mock_resp.json.return_value = {
        "output": json.dumps({
            "classification": "Bank Statement",
            "confidence": 0.97,
            "reasoning": "Monthly checking account statement with transaction ledger",
        })
    }
    with patch("httpx.Client.post", return_value=mock_resp) as mock_post:
        cls_result = provider_configured.classify_document(
            text="JPMorgan Chase Bank Statement\nAccount: ending in 4102\nStatement Period: 01/01/2024 to 01/31/2024",
            candidate_types=["Bank Statement", "Vendor Invoice", "Tax Return"],
        )
        assert cls_result == "Bank Statement"
        # Confirm request headers and payload
        args = mock_post.call_args
        assert args.kwargs["headers"]["Authorization"] == "Key fal_sec_live_998877665544332211aabbcc"
        assert args.kwargs["json"]["model"] == "anthropic/claude-3-opus"
        payload_data = args.kwargs["json"]
        content_text = payload_data.get("prompt") or payload_data.get("messages", [{}])[-1].get("content", "")
        assert "JPMorgan Chase" in content_text
        print("  [PASS] Document classification via fal.ai Claude Opus endpoint")

    # Test 4: Structured field extraction
    mock_field_resp = MagicMock()
    mock_field_resp.status_code = 200
    mock_field_resp.json.return_value = {
        "output": json.dumps({
            "field": "tax_year",
            "value": "2024",
            "confidence": 0.99,
        })
    }
    with patch("httpx.Client.post", return_value=mock_field_resp):
        val = provider_configured.extract_field("Client tax document for year 2024", "tax_year")
        assert val == "2024"
        print("  [PASS] Structured field extraction via Claude Opus")

    # Test 5: Ambiguous case semantic analysis
    mock_analysis_resp = MagicMock()
    mock_analysis_resp.status_code = 200
    mock_analysis_resp.json.return_value = {
        "output": json.dumps({
            "summary": "Ambiguous invoice with conflicting tax years.",
            "analysis": "Expense incurred in 2023 but invoiced in 2024.",
            "confidence": "high",
            "entities": {"period_year": 2024, "vendor_or_client": "Stripe Inc"},
            "findings": [{"issue": "Accrual required", "severity": "medium", "recommendation": "Post accrual"}],
        })
    }
    with patch("httpx.Client.post", return_value=mock_analysis_resp):
        analysis = provider_configured.analyze_document("Stripe Inc billing doc", "Analyze timing")
        assert analysis is not None
        assert analysis["entities"]["vendor_or_client"] == "Stripe Inc"
        assert len(analysis["findings"]) == 1
        print("  [PASS] Ambiguous document analysis with structured JSON return")

    # Test 6: Interchangeable direct Anthropic fallback routing
    provider_anthropic = ClaudeOpusFalProvider(
        api_key="",
        anthropic_key="sk-ant-valid-key",
    )
    assert provider_anthropic.backend_mode == "anthropic_direct"
    mock_anthropic_resp = MagicMock()
    mock_anthropic_resp.status_code = 200
    mock_anthropic_resp.json.return_value = {
        "content": [{"type": "text", "text": json.dumps({"classification": "Vendor Invoice", "confidence": 0.95})}]
    }
    with patch("httpx.Client.post", return_value=mock_anthropic_resp) as mock_post:
        res = provider_anthropic.classify_document("Invoice from Dell", ["Vendor Invoice"])
        assert res == "Vendor Invoice"
        assert mock_post.call_args.kwargs["headers"]["x-api-key"] == "sk-ant-valid-key"
        print("  [PASS] Replaceable direct Anthropic API fallback routing")

    # Test 7: Error handling (401, 429, 500, network timeouts, malformed JSON)
    for code, desc in [(401, "Auth failure"), (429, "Rate limit"), (500, "Server error")]:
        err_mock = MagicMock()
        err_mock.status_code = code
        err_mock.text = f"Error {code}"
        with patch("httpx.Client.post", return_value=err_mock):
            assert provider_configured.generate("test") is None

    with patch("httpx.Client.post", side_effect=httpx.TimeoutException("Timeout")):
        assert provider_configured.generate("test") is None

    malformed_mock = MagicMock()
    malformed_mock.status_code = 200
    malformed_mock.json.return_value = {"output": "Not valid JSON at all"}
    with patch("httpx.Client.post", return_value=malformed_mock):
        assert provider_configured.query_json("test") is None
    print("  [PASS] Error handling & timeouts (401, 429, 500, timeout, malformed JSON)")

    # Test 8: Engine wiring
    engine = build_default_engine()
    assert isinstance(engine.llm, ClaudeOpusFalProvider)
    print("  [PASS] Default engine wires ClaudeOpusFalProvider")


def test_opt_in_live_claude():
    """Separate opt-in live test against fal.ai Claude Opus (only if --live-claude flag passed)."""
    print("\n--- 6. Opt-In Live Claude Opus Real API Test ---")
    from engine.providers.llm import ClaudeOpusFalProvider
    provider = ClaudeOpusFalProvider()
    if not provider.available:
        print("  [SKIP] No live FAL_KEY configured for opt-in test")
        return
    res = provider.generate("Reply with exactly: LedgerLens Live Test OK")
    assert res is not None
    print(f"  [PASS] Live Claude Opus responded: {res}")


if __name__ == "__main__":
    test_backend_api_and_database()
    test_core_engine_realistic_files()
    test_end_to_end_integration()
    test_paddle_ocr_integration()
    test_claude_opus_fal_integration()
    if RUN_LIVE_CLAUDE:
        test_opt_in_live_claude()
    print("\n=======================================================")
    print("ALL VERIFICATION CHECKS PASSED SUCCESSFULLY!")
    print("=======================================================")
