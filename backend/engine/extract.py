"""Text extraction for supported file types.

Returns (text, status, reason). status is one of:
  ok               - text extracted
  empty            - file readable but no usable text
  needs_ocr        - scanned image / image-only PDF (OCR intentionally
                     deferred in this prototype; runs locally in desktop build)
  password         - encrypted / password protected
  error            - could not read / corrupted
"""
import csv
import io

IMAGE_EXTS = {"jpg", "jpeg", "png", "tiff", "tif"}


def extract_text(path: str, ext: str):
    ext = (ext or "").lower().lstrip(".")
    try:
        if ext == "pdf":
            return _extract_pdf(path)
        if ext == "docx":
            return _extract_docx(path)
        if ext == "xlsx":
            return _extract_xlsx(path)
        if ext == "csv":
            return _extract_csv(path)
        if ext in IMAGE_EXTS:
            return "", "needs_ocr", "Scanned image requires OCR"
        return "", "error", f"Unsupported file type: .{ext}"
    except Exception as e:  # noqa: BLE001 - engine must never crash the scan
        return "", "error", f"Could not read file: {type(e).__name__}"


def _extract_pdf(path: str):
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(path)
    except PdfReadError:
        return "", "error", "Corrupted or invalid PDF"

    if reader.is_encrypted:
        try:
            if reader.decrypt("") == 0:
                return "", "password", "Password-protected PDF"
        except Exception:  # noqa: BLE001
            return "", "password", "Password-protected PDF"

    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001
            continue
    text = "\n".join(parts).strip()
    if not text:
        return "", "needs_ocr", "Image-only / scanned PDF requires OCR"
    return text, "ok", ""


def _extract_docx(path: str):
    import docx

    doc = docx.Document(path)
    text = "\n".join(p.text for p in doc.paragraphs).strip()
    if not text:
        return "", "empty", "Document contains no text"
    return text, "ok", ""


def _extract_xlsx(path: str):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell is not None:
                    parts.append(str(cell))
    wb.close()
    text = " ".join(parts).strip()
    if not text:
        return "", "empty", "Spreadsheet contains no data"
    return text, "ok", ""


def _extract_csv(path: str):
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        sample = f.read(200000)
    rows = list(csv.reader(io.StringIO(sample)))
    text = " ".join(" ".join(c for c in r if c) for r in rows).strip()
    if not text:
        return "", "empty", "CSV contains no data"
    return text, "ok", ""
