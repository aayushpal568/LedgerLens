"""Report export helpers (CSV / XLSX / PDF).

Return raw bytes so the API can stream them. Framework-free.
"""
import csv
import io
from datetime import datetime
from xml.sax.saxutils import escape

CATEGORY_LABELS = {
    "exact_duplicate": "Exact Duplicate",
    "possible_duplicate": "Possible Duplicate",
    "missing_doc": "Missing Document",
    "wrong_period": "Wrong Period",
    "wrong_type": "Wrong Document Type",
    "unreadable": "Unreadable File",
}

STATUS_LABELS = {
    "unreviewed": "Unreviewed",
    "keep": "Keep",
    "keep_both": "Keep Both",
    "ignore": "Ignore",
    "review_later": "Review Later",
}

HEADERS = ["Category", "Issue", "Files", "Confidence", "Review Status", "Notes", "Evidence"]


def _spreadsheet_safe(value):
    """Prevent exported user-controlled text from becoming a spreadsheet formula."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _is_ai_involved(f: dict) -> tuple:
    """Check if AI / LLM contributed to this finding, its files, or its evidence."""
    tags = []
    if f.get("ai_assisted") or f.get("provenance") == "llm":
        tags.append("AI-Assisted")
    if f.get("classified_by") == "llm":
        tags.append("AI Classified")
    if f.get("detected_by") == "llm":
        tags.append("AI Period Detected")
    ev = f.get("evidence") or {}
    if ev.get("classified_by") == "llm":
        tags.append("AI Classified")
    if ev.get("detected_by") == "llm":
        tags.append("AI Period Detected")
    for file_info in f.get("files", []):
        if file_info.get("classified_by") == "llm":
            tags.append("AI Classified File")
        period = file_info.get("period") or {}
        if isinstance(period, dict) and period.get("detected_by") == "llm":
            tags.append("AI Period Detected File")
    if tags:
        return True, ", ".join(dict.fromkeys(tags))
    return False, ""


def _rows(findings):
    for f in findings:
        is_ai, ai_label = _is_ai_involved(f)
        conf_str = f"{f.get('confidence', 0)}% ({f.get('confidence_level', '')})"
        ev_str = (f.get("evidence", {}) or {}).get("summary", "")

        if is_ai:
            conf_str += f" [AI: {ai_label}]"
            ev_str = f"{ev_str} (AI/LLM: {ai_label})" if ev_str else f"AI/LLM: {ai_label}"

        yield [
            CATEGORY_LABELS.get(f["category"], f["category"]),
            f.get("title", ""),
            "; ".join(x["name"] for x in f.get("files", [])),
            conf_str,
            STATUS_LABELS.get(f.get("status", "unreviewed"), f.get("status", "")),
            f.get("note", "") or "",
            ev_str,
        ]


def to_csv(findings) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(HEADERS)
    for row in _rows(findings):
        w.writerow([_spreadsheet_safe(value) for value in row])
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(findings) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Exceptions"
    header_fill = PatternFill("solid", fgColor="0F766E")
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for row in _rows(findings):
        ws.append([_spreadsheet_safe(value) for value in row])
    widths = [20, 46, 34, 20, 16, 30, 50]
    for i, wdt in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = wdt
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(findings, meta: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer,
    )

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(letter), topMargin=0.5 * inch)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("t", parent=styles["Title"], textColor=colors.HexColor("#0F766E"))
    cell_style = ParagraphStyle("c", parent=styles["BodyText"], fontSize=8, leading=10)

    elems = [
        Paragraph("Document Review Report", title_style),
        Paragraph(
            f"Client: {escape(str(meta.get('client_name', '-')))} &nbsp;&nbsp; "
            f"Period: {escape(str(meta.get('expected_period', '-')))} &nbsp;&nbsp; "
            f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
            styles["Normal"],
        ),
        Spacer(1, 6),
        Paragraph(
            f"Total exceptions: {len(findings)}. Internal review artifact — no external actions taken.",
            styles["Italic"],
        ),
        Spacer(1, 12),
    ]

    data = [HEADERS]
    for row in _rows(findings):
        data.append([Paragraph(escape(str(c)), cell_style) for c in row])

    table = Table(data, repeatRows=1, colWidths=[70, 150, 130, 70, 60, 90, 150])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F766E")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    elems.append(table)
    doc.build(elems)
    return out.getvalue()
